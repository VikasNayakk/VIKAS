"""
OCR Notebook Image → Structured Excel Pipeline
Aliens Classes - Notebook-To-Excel-AI
"""
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_PATH = os.path.join(BASE, "input", "notebook_images", "Image.jpg")
OUTPUT_DIR = os.path.join(BASE, "output", "excel_files")
FINAL_XLSX = os.path.join(BASE, "output", "final_master.xlsx")
RAW_TEXT_DIR = os.path.join(BASE, "input", "raw_text")

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(RAW_TEXT_DIR, exist_ok=True)

# Step 1: OCR
print("[1/3] Running OCR on Image.jpg...")
try:
    import easyocr
    reader = easyocr.Reader(['en', 'hi'], gpu=False)
    results = reader.readtext(IMG_PATH, detail=0, paragraph=True)
    ocr_text = "\n".join(results)
except Exception as e:
    print(f"EasyOCR failed: {e}")
    # Fallback: try pytesseract
    try:
        import pytesseract
        from PIL import Image
        for p in [r'C:\Program Files\Tesseract-OCR\tesseract.exe',
                  r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe']:
            if os.path.exists(p):
                pytesseract.pytesseract.tesseract_cmd = p
                break
        img = Image.open(IMG_PATH)
        ocr_text = pytesseract.image_to_string(img)
    except Exception as e2:
        print(f"Pytesseract also failed: {e2}")
        print("Cannot OCR. Exiting.")
        sys.exit(1)

# Save raw text
raw_file = os.path.join(RAW_TEXT_DIR, "Image_ocr.txt")
with open(raw_file, "w", encoding="utf-8") as f:
    f.write(ocr_text)
print(f"   Raw text saved: {raw_file}")
print(f"   Extracted text preview:\n   {ocr_text[:500]}\n")

# Step 2: Parse and structure data
print("[2/3] Structuring data...")

import re
from datetime import datetime

lines = [l.strip() for l in ocr_text.split("\n") if l.strip()]

# Try to extract date
date_str = ""
for line in lines:
    # Common date patterns
    m = re.search(r'(\d{1,2})[/\-\.](\d{1,2})[/\-\.](\d{2,4})', line)
    if m:
        d, mo, y = m.group(1), m.group(2), m.group(3)
        if len(y) == 2:
            y = "20" + y
        date_str = f"{d.zfill(2)}-{mo.zfill(2)}-{y}"
        break

# Meeting name: use first meaningful line
meeting_name = lines[0] if lines else "Aliens Class"

# Type detection
full_text_lower = ocr_text.lower()
type_keywords = {
    "Development Team": ["coding", "code", "python", "ai", "backend", "software", "system", "develop", "program", "html", "css", "javascript", "react", "api", "database", "server"],
    "Graphics": ["ui", "ux", "design", "graphic", "photoshop", "figma", "visual", "asset", "canva", "illustrator"],
    "Marketing": ["marketing", "promotion", "ads", "growth", "campaign", "seo", "social media", "content"],
    "Sales": ["sales", "selling", "client", "revenue", "deal", "customer", "pricing"],
}

detected_type = "Common"
max_hits = 0
for t, keywords in type_keywords.items():
    hits = sum(1 for kw in keywords if kw in full_text_lower)
    if hits > max_hits:
        max_hits = hits
        detected_type = t

# Generate agenda (first 5-10 words summary from content)
content_words = " ".join(lines[1:5]) if len(lines) > 1 else meeting_name
words = content_words.split()
agenda = " ".join(words[:10]) if len(words) >= 5 else content_words
if len(agenda) > 80:
    agenda = agenda[:77] + "..."

row = {
    "Meeting Name": meeting_name,
    "Agenda": agenda,
    "Type": detected_type,
    "Duration": "",
    "Date": date_str,
    "Host": "A'dii Sunlay",
    "Present": "",
    "Late": "",
    "Absent": ""
}

print(f"   Meeting Name: {row['Meeting Name']}")
print(f"   Agenda: {row['Agenda']}")
print(f"   Type: {row['Type']}")
print(f"   Date: {row['Date']}")
print()

# Step 3: Write to Excel
print("[3/3] Writing to Excel...")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

columns = ["Meeting Name", "Agenda", "Type", "Duration", "Date", "Host", "Present", "Late", "Absent"]

# Check if final_master.xlsx exists and has data
if os.path.exists(FINAL_XLSX):
    try:
        wb = load_workbook(FINAL_XLSX)
        ws = wb.active
        # Check if headers exist
        if ws.max_row and ws.cell(1, 1).value:
            # Append row
            next_row = ws.max_row + 1
            for i, col in enumerate(columns, 1):
                ws.cell(next_row, i, row[col])
            print(f"   Appended to existing file at row {next_row}")
        else:
            raise ValueError("Empty sheet")
    except:
        wb = Workbook()
        ws = wb.active
        ws.title = "Aliens Classes"
        # Write headers
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for i, col in enumerate(columns, 1):
            cell = ws.cell(1, i, col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        # Write data row
        for i, col in enumerate(columns, 1):
            cell = ws.cell(2, i, row[col])
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        # Auto-width
        for i, col in enumerate(columns, 1):
            ws.column_dimensions[chr(64+i)].width = max(len(col) + 5, 18)
        print("   Created new Excel with headers + 1 data row")
else:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aliens Classes"
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for i, col in enumerate(columns, 1):
        cell = ws.cell(1, i, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    for i, col in enumerate(columns, 1):
        cell = ws.cell(2, i, row[col])
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[chr(64+i)].width = max(len(col) + 5, 18)
    print("   Created new Excel with headers + 1 data row")

wb.save(FINAL_XLSX)
print(f"\n✅ DONE! Excel saved: {FINAL_XLSX}")
print(f"   Raw OCR text: {raw_file}")

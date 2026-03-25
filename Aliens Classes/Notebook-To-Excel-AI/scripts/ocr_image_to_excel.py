"""
OCR notebook images → Extract class/meeting info → Append to SAME final_master.xlsx
Rules:
  - NEVER create new Excel, always append to existing final_master.xlsx
  - Image top heading = Meeting Name (class/meeting ka naam)
  - Read full page properly, extract date/topic/content
  - Host is always A'dii Sunlay
"""
import os
import re
import sys
import glob
import easyocr
from PIL import Image, ImageEnhance
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE = r"c:\Users\VikasNayak\OneDrive - Aliens Company\VIKAS\Aliens Classes\Notebook-To-Excel-AI"
IMG_DIR = os.path.join(BASE, "input", "notebook_images")
OUTPUT = os.path.join(BASE, "output", "final_master.xlsx")
RAW_TEXT_DIR = os.path.join(BASE, "input", "raw_text")
os.makedirs(RAW_TEXT_DIR, exist_ok=True)
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)

COLUMNS = ["Meeting Name", "Agenda", "Type", "Duration", "Date", "Host", "Present", "Late", "Absent"]

# ─── Styles ───
data_font = Font(name="Calibri", size=11)
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color="B0B0B0"), right=Side(style='thin', color="B0B0B0"),
    top=Side(style='thin', color="B0B0B0"), bottom=Side(style='thin', color="B0B0B0"),
)
alt_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
col_widths = {"A": 28, "B": 48, "C": 20, "D": 12, "E": 14, "F": 16, "G": 10, "H": 8, "I": 10}


def get_or_create_workbook():
    """Always use SAME Excel file. Create with headers only if missing."""
    if os.path.exists(OUTPUT):
        try:
            wb = load_workbook(OUTPUT)
            ws = wb.active
            if ws.cell(1, 1).value:
                return wb, ws
        except PermissionError:
            print("ERROR: Excel file is open! Please close it first.")
            sys.exit(1)
    # Create fresh with headers
    wb = Workbook()
    ws = wb.active
    ws.title = "Aliens Classes"
    for i, col in enumerate(COLUMNS, 1):
        cell = ws.cell(1, i, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[1].height = 30
    for c, w in col_widths.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"
    return wb, ws


def enhance_image(img_path):
    """Enhance image for better OCR: grayscale + contrast + sharpen + resize."""
    img = Image.open(img_path)
    img_gray = img.convert("L")
    img_gray = ImageEnhance.Contrast(img_gray).enhance(2.0)
    img_gray = ImageEnhance.Sharpness(img_gray).enhance(1.5)
    MAX_W = 2000
    if img_gray.width > MAX_W:
        ratio = MAX_W / img_gray.width
        img_gray = img_gray.resize((MAX_W, int(img_gray.height * ratio)), Image.Resampling.LANCZOS)
    temp = os.path.join(RAW_TEXT_DIR, "_temp_enhanced.jpg")
    img_gray.save(temp, quality=95)
    return temp


def ocr_image(img_path, reader):
    """Run OCR and return sorted text blocks with positions."""
    temp = enhance_image(img_path)
    results = reader.readtext(temp, detail=1, paragraph=False)
    try:
        os.remove(temp)
    except:
        pass
    # Sort by Y position (top to bottom), then X (left to right)
    results.sort(key=lambda x: (x[0][0][1], x[0][0][0]))
    return results


def extract_meeting_data(results, img_name):
    """Extract Meeting Name (top heading), Date, Type, Agenda from OCR results."""
    # Filter low confidence junk
    good_results = [(bbox, text, conf) for bbox, text, conf in results if conf > 0.3 and len(text.strip()) > 1]

    if not good_results:
        # Fallback: use all results
        good_results = [(bbox, text, conf) for bbox, text, conf in results if len(text.strip()) > 0]

    all_texts = [text.strip() for _, text, _ in good_results]
    full_text = " ".join(all_texts)

    # Meeting Name = first readable text block (top of page = heading)
    meeting_name = all_texts[0] if all_texts else img_name.replace(".jpg", "").replace(".png", "")

    # Date extraction
    date_str = ""
    for text in all_texts:
        m = re.search(r'(\d{1,2})[/\-\.](\d{1,2})[/\-\.](\d{2,4})', text)
        if m:
            d, mo, y = m.group(1), m.group(2), m.group(3)
            if len(y) == 2:
                y = "20" + y
            date_str = f"{d.zfill(2)}-{mo.zfill(2)}-{y}"
            break

    # Type detection
    full_lower = full_text.lower()
    type_map = {
        "Development Team": ["coding", "code", "python", "ai", "backend", "software", "system", "develop", "program",
                             "html", "css", "javascript", "react", "api", "database", "server", "git", "debug"],
        "Graphics": ["ui", "ux", "design", "graphic", "photoshop", "figma", "visual", "asset", "canva", "logo"],
        "Marketing": ["marketing", "promotion", "ads", "growth", "campaign", "seo", "social media", "branding"],
        "Sales": ["sales", "selling", "client", "revenue", "deal", "customer", "pricing", "lead"],
    }
    detected_type = "Common"
    max_hits = 0
    for t, keywords in type_map.items():
        hits = sum(1 for kw in keywords if kw in full_lower)
        if hits > max_hits:
            max_hits = hits
            detected_type = t

    # Agenda: summarize from content after heading (5-10 words)
    content = " ".join(all_texts[1:6]) if len(all_texts) > 1 else meeting_name
    words = content.split()
    agenda = " ".join(words[:10])
    if len(agenda) > 80:
        agenda = agenda[:77] + "..."
    if not agenda:
        agenda = meeting_name

    return {
        "Meeting Name": meeting_name,
        "Agenda": agenda,
        "Type": detected_type,
        "Duration": "",
        "Date": date_str,
        "Host": "A'dii Sunlay",
        "Present": "",
        "Late": "",
        "Absent": "",
    }


def append_to_excel(wb, ws, row_data):
    """Append one row to existing worksheet."""
    next_row = ws.max_row + 1
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=row_data[col_name])
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = left_align if col_name in ("Meeting Name", "Agenda") else center
    if next_row % 2 == 0:
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=next_row, column=col_idx).fill = alt_fill
    ws.row_dimensions[next_row].height = 22
    ws.auto_filter.ref = f"A1:I{next_row}"
    return next_row


# ─── MAIN ───
if __name__ == "__main__":
    # Find all images in notebook_images folder
    images = []
    for ext in ("*.jpg", "*.jpeg", "*.png", "*.bmp"):
        images.extend(glob.glob(os.path.join(IMG_DIR, ext)))

    if not images:
        print("No images found in input/notebook_images/")
        sys.exit(0)

    print(f"Found {len(images)} image(s) to process")
    print(f"Output: {OUTPUT}\n")

    # Init OCR reader once
    print("Loading OCR model...")
    reader = easyocr.Reader(["en"], gpu=False, verbose=False)

    wb, ws = get_or_create_workbook()

    for img_path in images:
        img_name = os.path.basename(img_path)
        print(f"\n--- Processing: {img_name} ---")

        # OCR
        results = ocr_image(img_path, reader)
        print(f"   Found {len(results)} text blocks")

        # Save raw OCR text
        raw_file = os.path.join(RAW_TEXT_DIR, img_name.rsplit(".", 1)[0] + "_ocr.txt")
        with open(raw_file, "w", encoding="utf-8") as f:
            for bbox, text, conf in results:
                f.write(f"[{conf:.2f}] {text}\n")

        # Extract structured data
        row_data = extract_meeting_data(results, img_name)

        print(f"   Meeting Name: {row_data['Meeting Name']}")
        print(f"   Agenda: {row_data['Agenda']}")
        print(f"   Type: {row_data['Type']}")
        print(f"   Date: {row_data['Date'] or '(not found)'}")

        # Append to Excel
        row_num = append_to_excel(wb, ws, row_data)
        print(f"   Added at row {row_num}")

    # Save SAME file
    wb.save(OUTPUT)
    total = ws.max_row - 1
    print(f"\n{'='*60}")
    print(f"DONE! {len(images)} image(s) processed")
    print(f"Total entries in Excel: {total}")
    print(f"Excel: {OUTPUT}")
    print(f"{'='*60}")

